"""EIA STEO world liquid-fuels balance → ``site/data/balance.json`` (ARCHITECTURE §2, §3.2; brief §5.4, §6.2).

Data path in a live run:

1. **EIA API v2** ``/v2/steo/data/`` — one request with all seven ``facets[seriesId][]``,
   ``start=2024-01``, ``length=5000``, key from ``env["EIA_API_KEY"]``. Without a key the API is not
   even attempted.
2. **Fallback** (no key, 403/429 or any other error): ``STEO_m.xlsx``, sheet ``3atab`` (no key needed),
   parsed with openpyxl in read-only mode. The xlsx carries no Brent/WTI rows, so those are ``null``.
3. The **quote** and the **release dates** come from ``global_oil.php`` (regex over the HTML). If that
   page fails: ``quote=null``, ``next_release=null`` and ``steo_release`` falls back to the xlsx
   "Forecast date:" cell or, failing that, to the first day of the current month
   (``info["release_date_estimated"]=True``).

Fixtures mode reads ``tests/fixtures/eia/steo_api.json`` (default) or ``steo_3atab.xlsx`` (when
``env["CRACKSPREAD_STEO_VIA"] == "xlsx"``) plus ``global_oil.html``; it never touches the network.

All diagnostics go through :func:`common.log` (stderr); nothing is printed to stdout and nothing calls
``sys.exit``. Runs on Python 3.9 and 3.12.
"""
from __future__ import annotations

import html as _html
import io
import json
import os
import re
from datetime import date, datetime
from typing import Any, Dict, List, Mapping, Optional, Tuple

import openpyxl

import common

SOURCE_KEY = "eia_steo"
OUTPUT_FILE = "balance.json"
SCHEMA = "balance"
STALE_KIND = "steo"

API_URL = "https://api.eia.gov/v2/steo/data/"
XLSX_URL = "https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx"
PAGE_URL = "https://www.eia.gov/outlooks/steo/report/global_oil.php"
SOURCE_NAME = "EIA Short-Term Energy Outlook (STEO)"
UNIT = "million barrels per day"
SHEET = "3atab"
MAX_XLSX_BYTES = 25 * 1024 * 1024   # STEO_m.xlsx is ~1.1 MB; anything far bigger is not the workbook
MAX_XLSX_ROWS = 2000                # 3atab has ~100 rows
START_PERIOD = "2024-01"
API_LENGTH = 5000
QUOTE_PHRASE = "global oil inventories"

#: (output field, API seriesId, xlsx column-A id, kind). Order = order of the fields in a month entry.
SERIES: Tuple[Tuple[str, str, Optional[str], str], ...] = (
    ("production", "PAPR_WORLD", "papr_world", "mbd"),
    ("consumption", "PATC_WORLD", "patc_world", "mbd"),
    ("stock_draw", "T3_STCHANGE_WORLD", "t3_stchange_world", "mbd"),
    ("opec", "PAPR_OPEC", "papr_opec", "mbd"),
    ("nonopec", "PAPR_NONOPEC", "papr_nonopec", "mbd"),
    ("brent", "BREPUUS", None, "price"),
    ("wti", "WTIPUUS", None, "price"),
)
FIELDS: Tuple[str, ...] = tuple(s[0] for s in SERIES)
API_SERIES_IDS: Tuple[str, ...] = tuple(s[1] for s in SERIES)
_API_TO_FIELD: Dict[str, str] = {s[1]: s[0] for s in SERIES}
_XLSX_TO_FIELD: Dict[str, str] = {s[2]: s[0] for s in SERIES if s[2]}
_KIND: Dict[str, str] = {s[0]: s[3] for s in SERIES}
CORE_FIELDS: Tuple[str, ...] = ("production", "consumption", "stock_draw")
PLAUSIBLE_MBD = (70.0, 130.0)

_MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
_MONTH_ABBR = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")
_MONTH_NAMES = ("January", "February", "March", "April", "May", "June", "July", "August",
                "September", "October", "November", "December")
_LONG_DATE = r"([A-Za-z]{3,9}\.?\s+\d{1,2},?\s+\d{4})"
_RELEASE_RE = re.compile(r"(?<!next )release date:?\s*" + _LONG_DATE, re.IGNORECASE)
_NEXT_RELEASE_RE = re.compile(r"next release date:?\s*" + _LONG_DATE, re.IGNORECASE)
_FORECAST_COMPLETED_RE = re.compile(r"forecast completed:?\s*" + _LONG_DATE, re.IGNORECASE)
_EDITION_RE = re.compile(r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{4})\b")

SeriesMap = Dict[str, Dict[str, Optional[float]]]


# ------------------------------------------------------------------------------- small helpers

def _to_float(value: Any) -> Optional[float]:
    """STEO value → float; ``None`` for missing markers (``None``, ``""``, ``"."``, ``"NA"``, ``"--"``, ``"w"``)."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return None if number != number else number
    text = str(value).strip().replace(",", "")
    if text in ("", ".", "-", "--", "NA", "N/A", "na", "n/a", "w", "W", "null"):
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return None if number != number else number


def _round(value: Optional[float], kind: str) -> Optional[float]:
    """Round for byte-stable output (mb/d and STEO monthly prices: 2 decimals). ``-0.0`` → ``0.0``."""
    if value is None:
        return None
    digits = 2 if kind in ("mbd", "price") else 3
    rounded = round(float(value), digits)
    return 0.0 if rounded == 0 else rounded


def _mean(values: List[Optional[float]]) -> Optional[float]:
    if not values or any(v is None for v in values):
        return None
    return sum(float(v) for v in values if v is not None) / len(values)


def _empty_series() -> SeriesMap:
    return {field: {} for field in FIELDS}


def _quarter_of(period: str) -> str:
    year, month = period.split("-")
    return "%sQ%d" % (year, (int(month) - 1) // 3 + 1)


def _edition_from_date(iso_date: Optional[str]) -> Optional[str]:
    """``"2026-10-06"`` → ``"October 2026"`` (STEO editions are named after their release month)."""
    if not iso_date:
        return None
    try:
        d = common.parse_iso(iso_date)
    except ValueError:
        return None
    return "%s %d" % (_MONTH_NAMES[d.month - 1], d.year)


def _parse_long_date(text: Any) -> Optional[str]:
    """``"Thursday, October 1, 2026"`` / ``"October 6, 2026"`` / ``"Oct 6, 2026"`` / a date cell → ``"YYYY-MM-DD"``."""
    if isinstance(text, datetime):
        return text.date().isoformat()
    if isinstance(text, date):
        return text.isoformat()
    if not isinstance(text, str):
        return None
    cleaned = " ".join(text.replace("\xa0", " ").split()).strip(" :")
    if not re.search(r"\b(19|20)\d{2}\b", cleaned):
        return None  # refuse to guess a year
    for fmt in ("%B %d, %Y", "%b %d, %Y", "%b. %d, %Y", "%B %d %Y", "%A, %B %d, %Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(cleaned, fmt).date().isoformat()
        except ValueError:
            continue
    from dateutil import parser as _dateutil_parser  # lazy: unusual spellings only

    try:
        return _dateutil_parser.parse(cleaned, fuzzy=True).date().isoformat()
    except (ValueError, OverflowError):
        return None


def _redact(text: str, secret: str) -> str:
    text = common._redact(text) if hasattr(common, "_redact") else text
    return text.replace(secret, "***") if secret else text


def _response_text(resp: Any) -> str:
    """Decode a response as UTF-8 (EIA declares utf-8); fall back to the detected encoding."""
    content = getattr(resp, "content", None)
    if isinstance(content, (bytes, bytearray)):
        raw = bytes(content)
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            enc = getattr(resp, "apparent_encoding", None) or getattr(resp, "encoding", None) or "latin-1"
            return raw.decode(enc, errors="replace")
    return str(getattr(resp, "text", "") or "")


# ---------------------------------------------------------------------------------- API path

def api_params(key: str) -> List[Tuple[str, Any]]:
    """Query parameters of the single STEO request (repeated ``facets[seriesId][]`` keys)."""
    params: List[Tuple[str, Any]] = [("frequency", "monthly"), ("data[0]", "value")]
    params.extend(("facets[seriesId][]", sid) for sid in API_SERIES_IDS)
    params.extend([("start", START_PERIOD), ("length", API_LENGTH), ("api_key", key)])
    return params


def parse_api_payload(payload: Any) -> SeriesMap:
    """``response.data`` rows → ``{field: {period: float|None}}``. Values are strings in the API."""
    if not isinstance(payload, dict):
        raise ValueError("EIA API: unexpected payload type %s" % type(payload).__name__)
    if "response" not in payload:
        err = payload.get("error") or payload.get("message") or "no 'response' object"
        raise common.FetchError("EIA API error: %s" % str(err)[:200])
    response = payload.get("response")
    data = response.get("data") if isinstance(response, dict) else None
    if not isinstance(data, list):
        raise ValueError("EIA API: 'response.data' is not a list")
    try:
        total = int(response.get("total") or 0)
    except (TypeError, ValueError):
        total = 0
    if total > len(data):
        common.log("balance: warning: EIA API returned %d of %d rows (truncated)" % (len(data), total))
    series = _empty_series()
    for row in data:
        if not isinstance(row, dict):
            continue
        field = _API_TO_FIELD.get(str(row.get("seriesId", "")).strip().upper())
        period = str(row.get("period", "")).strip()
        if field is None or not _MONTH_RE.match(period):
            continue
        series[field][period] = _to_float(row.get("value"))
    if not any(series[f] for f in CORE_FIELDS):
        raise ValueError("EIA API: no usable STEO rows in the response")
    return series


def fetch_api(key: str, session: Any = None) -> SeriesMap:
    resp = common.http_get(API_URL, params=api_params(key), session=session)
    return parse_api_payload(resp.json())


# --------------------------------------------------------------------------------- xlsx path

def _cell_year(cell: Any) -> Optional[int]:
    if isinstance(cell, bool) or cell is None:
        return None
    if isinstance(cell, (int, float)):
        if float(cell).is_integer() and 1990 <= int(cell) <= 2100:
            return int(cell)
        return None
    text = str(cell).strip()
    return int(text) if re.fullmatch(r"(19|20)\d{2}", text) else None


def _cell(row: Tuple[Any, ...], index: int) -> Any:
    return row[index] if 0 <= index < len(row) else None


def parse_3atab_rows(rows: List[Tuple[Any, ...]]) -> Dict[str, Any]:
    """Parse the rows of sheet ``3atab`` (values only).

    Returns ``{"series": SeriesMap, "forecast_date": "YYYY-MM-DD"|None, "edition": "October 2026"|None}``.
    Layout: a row of month names (``Jan`` … ``Dec`` repeated) with the year in the row above it at each
    January column; data rows carry the series id in column A (first matching row wins).
    """
    header_idx: Optional[int] = None
    for i, row in enumerate(rows):
        count = sum(1 for c in row if isinstance(c, str) and c.strip().lower()[:3] in _MONTH_ABBR and len(c.strip()) <= 9)
        if count >= 12:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("%s: month header row not found" % SHEET)
    year_row: Tuple[Any, ...] = rows[header_idx - 1] if header_idx > 0 else ()
    col_period: Dict[int, str] = {}
    year: Optional[int] = None
    for j, cell in enumerate(rows[header_idx]):
        if not isinstance(cell, str):
            continue
        mon = cell.strip().lower()[:3]
        if mon not in _MONTH_ABBR or len(cell.strip()) > 9:
            continue
        found_year = _cell_year(_cell(year_row, j))
        if found_year is not None:
            year = found_year
        elif mon == "jan" and year is not None:
            year += 1  # tolerate a missing year label: a new January means a new year
        if year is None:
            continue
        col_period[j] = "%04d-%02d" % (year, _MONTH_ABBR.index(mon) + 1)
    if not col_period:
        raise ValueError("%s: no year labels above the month header" % SHEET)

    series = _empty_series()
    seen = set()
    for row in rows[header_idx + 1:]:
        head = _cell(row, 0)
        if not isinstance(head, str):
            continue
        field = _XLSX_TO_FIELD.get(head.strip().lower())
        if field is None or field in seen:
            continue
        seen.add(field)
        for j, period in col_period.items():
            series[field][period] = _to_float(_cell(row, j))
    missing_core = [sid for f, _api, sid, _k in SERIES if f in CORE_FIELDS and f not in seen]
    if missing_core:
        raise ValueError("%s: rows not found: %s" % (SHEET, ", ".join(missing_core)))
    missing_other = [sid for f, _api, sid, _k in SERIES if sid and f not in seen and f not in CORE_FIELDS]
    if missing_other:
        common.log("balance: warning: %s rows missing in %s (left null)" % (", ".join(missing_other), SHEET))

    forecast_date: Optional[str] = None
    edition: Optional[str] = None
    for i, row in enumerate(rows[: header_idx + 1]):
        for j, cell in enumerate(row):
            if not isinstance(cell, str):
                continue
            text = cell.strip()
            low = text.lower()
            if forecast_date is None and low.startswith("forecast date"):
                candidates: List[Any] = [text.split(":", 1)[1] if ":" in text else ""]
                candidates.extend(_cell(row, k) for k in range(j + 1, min(j + 4, len(row))))
                if i + 1 < len(rows):
                    candidates.append(_cell(rows[i + 1], j))
                for cand in candidates:
                    forecast_date = _parse_long_date(cand)
                    if forecast_date:
                        break
            if edition is None and "short-term energy outlook" in low:
                m = _EDITION_RE.search(text)
                if m:
                    edition = "%s %s" % (m.group(1), m.group(2))
    return {"series": series, "forecast_date": forecast_date, "edition": edition}


def parse_xlsx(source: Any) -> Dict[str, Any]:
    """``source`` = path or binary file-like of ``STEO_m.xlsx`` (or the trimmed fixture)."""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise ValueError("STEO workbook has no sheet %r (sheets: %s)" % (SHEET, ", ".join(wb.sheetnames[:8])))
        # 3atab has ~100 rows; the cap keeps a bloated or crafted sheet from exhausting memory.
        rows = [tuple(r) for r in wb[SHEET].iter_rows(values_only=True, max_row=MAX_XLSX_ROWS)]
    finally:
        wb.close()
    return parse_3atab_rows(rows)


def fetch_xlsx(session: Any = None) -> Dict[str, Any]:
    resp = common.http_get(XLSX_URL, session=session, max_bytes=MAX_XLSX_BYTES)
    content = bytes(resp.content)
    if not content.startswith(b"PK"):
        raise ValueError("STEO_m.xlsx: response is not an xlsx file (%d bytes)" % len(content))
    return parse_xlsx(io.BytesIO(content))


# --------------------------------------------------------------------------- global_oil.php

def html_to_lines(html: str) -> List[str]:
    """Visible text of an HTML document as whitespace-normalised lines (one per block element)."""
    s = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
    s = re.sub(r"(?s)<!--.*?-->", " ", s)
    # Block-level tags are the only paragraph boundaries; source line breaks inside a <p> are plain whitespace.
    s = re.sub(r"(?i)<\s*/?\s*(br|p|div|h[1-6]|li|tr|table|ul|ol|section|article|header|footer)\b[^>]*>", "\x00", s)
    s = re.sub(r"(?s)<[^>]+>", " ", s)
    s = _html.unescape(s).replace("\xa0", " ")
    lines = [" ".join(chunk.split()) for chunk in s.split("\x00")]
    return [line for line in lines if line]


def _sentence_at(text: str, idx: int) -> str:
    """The sentence of ``text`` that contains position ``idx`` (``U.S.``-style abbreviations tolerated)."""
    start = 0
    for m in re.finditer(r"(?<![ .][A-Z])[.!?]['\")\]]*\s+", text[:idx]):
        start = m.end()
    m = re.search(r"(?<![ .][A-Z])[.!?]['\")\]]*(?=\s|$)", text[idx:])
    end = idx + m.end() if m else len(text)
    return text[start:end].strip()


def extract_quote(html: str, phrase: str = QUOTE_PHRASE) -> Optional[str]:
    """First sentence of the page containing ``phrase`` (case-insensitive), or ``None``."""
    needle = phrase.lower()
    for line in html_to_lines(html):
        idx = line.lower().find(needle)
        if idx >= 0:
            return _sentence_at(line, idx)
    return None


def parse_page(html: str) -> Dict[str, Optional[str]]:
    """Quote + ``Release Date`` / ``Next Release Date`` / ``Forecast Completed`` from global_oil.php."""
    flat = " ".join(html_to_lines(html))
    out: Dict[str, Optional[str]] = {"quote": extract_quote(html), "release_date": None, "next_release": None,
                                     "forecast_completed": None}
    for key, rx in (("release_date", _RELEASE_RE), ("next_release", _NEXT_RELEASE_RE),
                    ("forecast_completed", _FORECAST_COMPLETED_RE)):
        m = rx.search(flat)
        if m:
            out[key] = _parse_long_date(m.group(1))
    return out


# --------------------------------------------------------------------------------- document

def classify_status(stock_draw: Optional[float], cfg: Mapping[str, Any]) -> str:
    """``deficit`` if the draw exceeds ``cfg.balance_thresholds.deficit`` (default 0.3), ``surplus`` if it
    is below ``…surplus`` (default -0.3), else ``balanced``."""
    thresholds = cfg.get("balance_thresholds") if isinstance(cfg, Mapping) else None
    thresholds = thresholds if isinstance(thresholds, Mapping) else {}
    deficit = float(thresholds.get("deficit", 0.3))
    surplus = float(thresholds.get("surplus", -0.3))
    if stock_draw is None:
        raise ValueError("balance: current month has no stock change value")
    value = float(stock_draw)
    if value > deficit:
        return "deficit"
    if value < surplus:
        return "surplus"
    return "balanced"


def build_doc(
    series: SeriesMap,
    cfg: Mapping[str, Any],
    now: datetime,
    *,
    via: str,
    page: Optional[Mapping[str, Any]] = None,
    xlsx_meta: Optional[Mapping[str, Any]] = None,
    fetched_at: Optional[str] = None,
) -> Tuple[dict, dict]:
    """Assemble ``balance.json`` from parsed series (pure; no I/O). Returns ``(doc, info)``.

    Raises ``PlausibilityError`` for production/consumption outside 70–130 mb/d and ``ValueError`` when
    the data does not cover the current month.
    """
    now = common.parse_iso(now)
    stamp = fetched_at or common.iso_utc(now)
    current_month = now.strftime("%Y-%m")
    periods = sorted({p for f in FIELDS for p in series.get(f, {}) if _MONTH_RE.match(p) and p >= START_PERIOD})
    if not periods:
        raise ValueError("balance: no STEO months from %s onward" % START_PERIOD)

    months: List[dict] = []
    for period in periods:
        entry: Dict[str, Any] = {"period": period}
        for field in FIELDS:
            entry[field] = _round(series.get(field, {}).get(period), _KIND[field])
        entry["is_forecast"] = period >= current_month
        months.append(entry)
        for field in ("production", "consumption"):
            if entry[field] is not None:
                common.ensure_range(entry[field], PLAUSIBLE_MBD[0], PLAUSIBLE_MBD[1], "%s %s" % (field, period))

    quarters: List[dict] = []
    by_quarter: Dict[str, List[str]] = {}
    for period in periods:
        by_quarter.setdefault(_quarter_of(period), []).append(period)
    for quarter in sorted(by_quarter):
        members = by_quarter[quarter]
        if len(members) != 3:
            continue  # means of complete quarters only
        q: Dict[str, Any] = {"period": quarter}
        for field in CORE_FIELDS:
            q[field] = _round(_mean([series.get(field, {}).get(p) for p in members]), "mbd")
        q["is_forecast"] = any(p >= current_month for p in members)
        quarters.append(q)

    current = next((m for m in months if m["period"] == current_month), None)
    if current is None:
        raise ValueError("balance: STEO data (%s…%s) has no entry for the current month %s"
                         % (periods[0], periods[-1], current_month))
    status = classify_status(current["stock_draw"], cfg)

    page = page or {}
    xlsx_meta = xlsx_meta or {}
    release = page.get("release_date")
    estimated = False
    if not release:
        release = xlsx_meta.get("forecast_date")
    if not release:
        release = now.strftime("%Y-%m-01")
        estimated = True
    edition = xlsx_meta.get("edition") or (None if estimated else _edition_from_date(release))
    quote = page.get("quote")

    doc: Dict[str, Any] = {
        "schema_version": 1,
        "generated_at": stamp,
        "as_of": release,
        "fetched_at": stamp,
        "stale": False,
        "source": "%s, %s" % (SOURCE_NAME, edition) if edition else SOURCE_NAME,
        "source_url": PAGE_URL,
        "data_url": API_URL if via == "api" else XLSX_URL,
        "via": via,
        "steo_release": release,
        "release_date_estimated": estimated,   # True = first of the month guessed (brief §7: never an unflagged estimate)
        "next_release": page.get("next_release"),
        "forecast_completed": page.get("forecast_completed"),
        "steo_edition": edition,
        "current_month": current_month,
        "months": months,
        "quarters": quarters,
        "current": current,
        "quote": quote,
        "quote_url": PAGE_URL,
        "status": status,
        "unit": UNIT,
        "series_ids": {field: sid for field, sid, _x, _k in SERIES},
        "notes": {
            "stock_draw": "Net inventory withdrawals (EIA T3_STCHANGE_WORLD): positive = inventories drawn down.",
            "is_forecast": "Months from the current month onward are EIA estimates/forecasts.",
            "prices": "Brent/WTI are STEO monthly averages in USD/bbl (API path only).",
        },
    }
    info: Dict[str, Any] = {
        "via": via,
        "release_date_estimated": estimated,
        "quote_found": quote is not None,
        "months": len(months),
    }
    return doc, info


# -------------------------------------------------------------------------------------- run

def _load_series(fixtures: bool, env: Mapping[str, str], session: Any) -> Tuple[SeriesMap, str, Dict[str, Any], Dict[str, Any]]:
    """→ ``(series, via, xlsx_meta, extra_info)``."""
    if fixtures:
        if (env.get("CRACKSPREAD_STEO_VIA") or "").strip().lower() == "xlsx":
            parsed = parse_xlsx(common.FIXTURES / "eia" / "steo_3atab.xlsx")
            return parsed["series"], "xlsx", parsed, {}
        with open(common.FIXTURES / "eia" / "steo_api.json", "r", encoding="utf-8") as fh:
            payload = json.load(fh)
        return parse_api_payload(payload), "api", {}, {}

    extra: Dict[str, Any] = {}
    key = (env.get("EIA_API_KEY") or "").strip()
    if not key:
        common.log("balance: EIA_API_KEY not set; using STEO_m.xlsx")
    else:
        try:
            return fetch_api(key, session), "api", {}, {}
        except Exception as exc:  # noqa: BLE001  (any failure → xlsx fallback)
            message = _redact("%s: %s" % (type(exc).__name__, exc), key)[:300]
            common.log("balance: EIA API failed (%s); falling back to STEO_m.xlsx" % message)
            extra["api_error"] = message
    parsed = fetch_xlsx(session)
    return parsed["series"], "xlsx", parsed, extra


def _load_page(fixtures: bool, session: Any) -> Tuple[Optional[Dict[str, Optional[str]]], Optional[str]]:
    """→ ``(parsed page or None, error text or None)``."""
    try:
        if fixtures:
            with open(common.FIXTURES / "eia" / "global_oil.html", "r", encoding="utf-8") as fh:
                html = fh.read()
        else:
            html = _response_text(common.http_get(PAGE_URL, session=session))
        parsed = parse_page(html)
    except Exception as exc:  # noqa: BLE001  (the page is optional: quote=null, dates fall back)
        message = ("%s: %s" % (type(exc).__name__, exc))[:300]
        common.log("balance: global_oil.php failed (%s); quote=null" % message)
        return None, message
    if parsed.get("quote") is None:
        common.log("balance: warning: no sentence with %r found on global_oil.php" % QUOTE_PHRASE)
    if parsed.get("release_date") is None:
        common.log("balance: warning: no 'Release Date:' found on global_oil.php")
    return parsed, None


def run(
    cfg: dict,
    old: Optional[dict],
    *,
    fixtures: bool = False,
    now: Optional[datetime] = None,
    session: Any = None,
    env: Optional[Mapping[str, str]] = None,
    context: Optional[dict] = None,
) -> Tuple[dict, dict]:
    """Build ``(doc, info)`` for ``balance.json`` (ARCHITECTURE §2). Raises when nothing usable was fetched."""
    del old, context  # the whole document is rebuilt every run; nothing is carried forward
    now_dt = common.now_utc() if now is None else common.parse_iso(now)
    env_map: Mapping[str, str] = os.environ if env is None else env
    stamp = common.iso_utc(now_dt)

    series, via, xlsx_meta, extra = _load_series(fixtures, env_map, session)
    page, page_error = _load_page(fixtures, session)
    doc, info = build_doc(series, cfg, now_dt, via=via, page=page, xlsx_meta=xlsx_meta, fetched_at=stamp)
    info.update(extra)
    if page_error:
        info["page_error"] = page_error
    common.log("balance: %s via %s, %s, current=%s status=%s%s" % (
        doc["steo_edition"] or "STEO", via, doc["steo_release"], doc["current_month"], doc["status"],
        " (release date estimated)" if info["release_date_estimated"] else ""))
    return doc, info
